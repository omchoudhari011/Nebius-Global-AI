"""The PC actions the server's model can ask for. Each takes the validated arguments and
returns a JSON-able result, or raises with a message the model reads.

Safety lives here as well as on the server: anything marked `confirm` was already approved
by you in a dialog before it gets here; files inside the protected folders (Windows, Program
Files, ...) are never changed; deleting goes to the Recycle Bin; nothing is overwritten
unless the request says so.
"""

from __future__ import annotations

import base64
import ctypes
import datetime as dt
import fnmatch
import glob
import io
import logging
import os
import re
import shutil
import time
import webbrowser
from pathlib import Path

from . import browser, notices

log = logging.getLogger(__name__)
MAX_READ = 8000
MAX_LIST = 40  # a long list costs the model seconds to write out; it says how many more
SKIP_DIRS = {"AppData", "node_modules", ".git", "$Recycle.Bin", "Windows", ".venv", "venv", "__pycache__", "site-packages",
             "anaconda3", "miniconda3", "envs", ".cache", ".conda", "Program Files", "Program Files (x86)", "ProgramData"}

# Common names people say, mapped to what Windows can start.
APPS = {
    "notepad": "notepad", "calculator": "calc", "calc": "calc", "paint": "mspaint", "explorer": "explorer",
    "file explorer": "explorer", "files": "explorer", "task manager": "taskmgr", "settings": "ms-settings:",
    "chrome": "chrome", "google chrome": "chrome", "edge": "msedge", "microsoft edge": "msedge", "firefox": "firefox",
    "word": "winword", "excel": "excel", "powerpoint": "powerpnt", "outlook": "outlook", "vs code": "code",
    "vscode": "code", "visual studio code": "code", "spotify": "spotify:", "camera": "microsoft.windows.camera:",
    "clock": "ms-clock:", "mail": "outlookmail:", "store": "ms-windows-store:", "terminal": "wt", "cmd": "cmd",
}
# Sites people name without ".com" ("open youtube").
SITES = {
    "youtube": "https://www.youtube.com", "gmail": "https://mail.google.com", "google": "https://www.google.com",
    "wikipedia": "https://www.wikipedia.org", "whatsapp web": "https://web.whatsapp.com", "instagram": "https://www.instagram.com",
    "facebook": "https://www.facebook.com", "twitter": "https://x.com", "x": "https://x.com", "linkedin": "https://www.linkedin.com",
    "netflix": "https://www.netflix.com", "amazon": "https://www.amazon.in", "flipkart": "https://www.flipkart.com",
    "github": "https://github.com", "chatgpt": "https://chatgpt.com", "google maps": "https://maps.google.com",
    "maps": "https://maps.google.com", "google drive": "https://drive.google.com", "drive": "https://drive.google.com",
    "google docs": "https://docs.google.com", "google calendar": "https://calendar.google.com", "reddit": "https://www.reddit.com",
    "outlook web": "https://outlook.live.com", "hotstar": "https://www.hotstar.com", "prime video": "https://www.primevideo.com",
    "youtube music": "https://music.youtube.com", "spotify web": "https://open.spotify.com", "stack overflow": "https://stackoverflow.com",
}
SEARCH_URLS = {
    "google": "https://www.google.com/search?q={}", "youtube": "https://www.youtube.com/results?search_query={}",
    "amazon": "https://www.amazon.in/s?k={}", "flipkart": "https://www.flipkart.com/search?q={}",
    "wikipedia": "https://en.wikipedia.org/w/index.php?search={}", "github": "https://github.com/search?q={}",
    "maps": "https://www.google.com/maps/search/{}", "images": "https://www.google.com/search?tbm=isch&q={}",
    "news": "https://news.google.com/search?q={}", "reddit": "https://www.reddit.com/search/?q={}",
    "spotify": "https://open.spotify.com/search/{}", "gmail": "https://mail.google.com/mail/u/0/#search/{}",
    "x": "https://x.com/search?q={}",
}
SETTINGS_PAGES = {
    "wifi": "network-wifi", "bluetooth": "bluetooth", "display": "display", "sound": "sound", "battery": "batterysaver",
    "apps": "appsfeatures", "update": "windowsupdate", "notifications": "notifications", "network": "network-status",
    "personalization": "personalization", "storage": "storagesense", "mouse": "mousetouchpad", "keyboard": "typing",
    "time": "dateandtime", "privacy": "privacy", "home": "",
}
# Never closed by close_app: Windows itself, and this app.
SYSTEM_PROCESSES = {"system", "csrss.exe", "wininit.exe", "winlogon.exe", "services.exe", "lsass.exe", "svchost.exe",
                    "smss.exe", "dwm.exe", "explorer.exe", "fontdrvhost.exe", "sihost.exe", "ctfmon.exe", "python.exe",
                    "pythonw.exe", "pythonw3.11.exe", "python3.11.exe"}
FILE_TYPES = {
    "Images": {"jpg", "jpeg", "png", "gif", "bmp", "webp", "heic", "svg", "tif", "tiff"},
    "Documents": {"pdf", "doc", "docx", "txt", "rtf", "odt", "md", "xls", "xlsx", "csv", "ppt", "pptx", "epub"},
    "Videos": {"mp4", "mkv", "mov", "avi", "wmv", "webm"},
    "Music": {"mp3", "wav", "flac", "m4a", "aac", "ogg"},
    "Archives": {"zip", "rar", "7z", "tar", "gz"},
    "Installers": {"exe", "msi"},
    "Code": {"py", "js", "ts", "html", "css", "json", "ipynb", "java", "c", "cpp"},
}
on_timer = None  # set by the app: called with the label when a timer rings
FILE_EXTENSIONS = {"pdf", "txt", "doc", "docx", "xls", "xlsx", "ppt", "pptx", "png", "jpg", "jpeg", "gif", "mp3", "mp4",
                   "wav", "zip", "exe", "msi", "md", "py", "js", "ts", "csv", "json", "html", "htm", "log", "lnk", "bat", "ps1"}
VK = {"play_pause": 0xB3, "next": 0xB0, "previous": 0xB1, "mute": 0xAD, "volume_up": 0xAF, "volume_down": 0xAE}


class ActionError(RuntimeError):
    pass


def _path(raw: str) -> Path:
    return Path(os.path.expandvars(os.path.expanduser(raw.strip().strip('"')))).resolve()


def _guard(path: Path, protected: list[str]) -> None:
    for root in protected:
        root_path = Path(root).resolve()
        if path == root_path or root_path in path.parents:
            raise ActionError(f"{path} is inside {root_path}, which Custodian never changes")


def _key(vk: int, times: int = 1) -> None:
    user32 = ctypes.windll.user32
    for _ in range(times):
        user32.keybd_event(vk, 0, 0, 0)
        user32.keybd_event(vk, 0, 2, 0)  # KEYEVENTF_KEYUP
        time.sleep(0.02)


# Looking ---------------------------------------------------------------------------------------

def system_info(**_: object) -> dict:
    """The time and the PC's vital signs: processor, memory, graphics, battery, drives, network,
    uptime and the busiest apps (stats.py)."""
    from .stats import SAMPLER

    if SAMPLER.fresh(30):  # the HUD read it moments ago: one more reading measures since then
        reading = SAMPLER.sample(with_apps=True)
    else:  # prime the counters, then measure over half a second: the apps busy right now
        SAMPLER.sample(with_apps=True)
        reading = SAMPLER.sample(cpu_interval=0.5, with_apps=True)
    return {"time": dt.datetime.now().strftime("%A %d %B %Y, %H:%M"), **reading, "active_window": _active_title()}


def _active_title() -> str | None:
    try:
        import pygetwindow

        window = pygetwindow.getActiveWindow()
        return window.title if window else None
    except Exception:  # noqa: BLE001
        return None


def list_files(path: str = "~", pattern: str = "*", **_: object) -> dict:
    folder = _path(path)
    if not folder.is_dir():
        raise ActionError(f"{folder} isn't a folder")
    items = []
    for entry in sorted(folder.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower())):
        if not fnmatch.fnmatch(entry.name.lower(), pattern.lower()):
            continue
        try:
            stat = entry.stat()
        except OSError:
            continue
        items.append({"name": entry.name, "type": "folder" if entry.is_dir() else "file",
                      "size_kb": None if entry.is_dir() else round(stat.st_size / 1024, 1),
                      "modified": dt.datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M")})
        if len(items) >= MAX_LIST:
            break
    return {"folder": str(folder), "items": items, "more": len(items) >= MAX_LIST}


def find_files(query: str, folder: str = "~", **_: object) -> dict:
    root, needle, found = _path(folder), query.lower(), []
    deadline = time.monotonic() + 10
    for current, dirs, files in os.walk(root):
        dirs[:] = [d for d in dirs if d not in SKIP_DIRS and not d.startswith(".")]
        for name in files + dirs:
            if needle in name.lower():
                found.append(str(Path(current) / name))
                if len(found) >= 30:
                    return {"query": query, "matches": found, "more": True}
        if time.monotonic() > deadline:
            return {"query": query, "matches": found, "more": True, "note": "stopped after 10 seconds"}
    return {"query": query, "matches": found, "more": False}


def read_file(path: str, **_: object) -> dict:
    file = _path(path)
    if not file.is_file():
        raise ActionError(f"{file} isn't a file")
    data = file.read_bytes()[: MAX_READ * 4]
    if b"\x00" in data[:2048]:
        raise ActionError(f"{file.name} isn't a text file")
    text = data.decode("utf-8", "replace")
    return {"path": str(file), "text": text[:MAX_READ] + (" …(cut)" if len(text) > MAX_READ else "")}


def list_windows(**_: object) -> dict:
    import pygetwindow

    return {"windows": [t for t in pygetwindow.getAllTitles() if t.strip()][:25], "active": _active_title()}


def get_clipboard(**_: object) -> dict:
    import pyperclip

    text = pyperclip.paste() or ""
    return {"text": text[:MAX_READ]}


def screenshot(**_: object) -> dict:
    import pyautogui

    image = pyautogui.screenshot()
    folder = Path.home() / "Pictures" / "Custodian"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"screenshot-{dt.datetime.now():%Y%m%d-%H%M%S}.png"
    image.save(path)
    small = image.copy()
    small.thumbnail((1280, 1280))
    buffer = io.BytesIO()
    small.save(buffer, format="PNG")
    return {"path": str(path), "image_b64": base64.b64encode(buffer.getvalue()).decode()}


_last_capture = None  # the screenshot capture() took, to compare with the one after a click


def _shot_b64(image) -> tuple[str, int, int]:  # noqa: ANN001 - a PIL image
    small = image.copy()
    small.thumbnail((1280, 1280))
    buffer = io.BytesIO()
    small.save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode(), small.width, small.height


def capture(purpose: str = "", **_: object) -> dict:
    """A screenshot for the vision model to read (pc_look, pc_click_on): sent, never saved here."""
    import pyautogui

    global _last_capture
    image = pyautogui.screenshot()
    _last_capture = image
    data, width, height = _shot_b64(image)
    return {"image_b64": data, "width": width, "height": height, "window": _active_title()}


def changed_share(before, after) -> float:  # noqa: ANN001 - PIL images
    """How much of the screen looks different, 0 to 1, on small greyscale copies (cheap, and
    blind to a blinking cursor)."""
    import numpy as np

    size = (160, 90)
    a, b = (np.asarray(image.convert("L").resize(size), dtype=np.int16) for image in (before, after))
    return float(np.mean(np.abs(a - b) > 24))


def click_at_checked(x: int, y: int, image_width: int, image_height: int, button: str = "left", **_: object) -> dict:
    """Click where the vision model found the target (in the screenshot's pixels), wait, and
    capture the screen again so the server can check it worked."""
    import pyautogui

    screen_w, screen_h = pyautogui.size()
    real_x, real_y = round(x * screen_w / image_width), round(y * screen_h / image_height)
    if not (0 <= real_x < screen_w and 0 <= real_y < screen_h):
        raise ActionError("that point is off the screen; nothing was clicked")
    before = _last_capture or pyautogui.screenshot()
    if button == "double":
        pyautogui.doubleClick(real_x, real_y)
    else:
        pyautogui.click(real_x, real_y, button=button)
    time.sleep(0.9)  # let the page or window react
    after = pyautogui.screenshot()
    data, _, _ = _shot_b64(after)
    return {"clicked": [real_x, real_y], "changed": round(changed_share(before, after), 4), "image_b64": data}


# Doing -----------------------------------------------------------------------------------------

def _start_menu_shortcut(name: str) -> Path | None:
    roots = [Path(os.environ.get("ProgramData", r"C:\ProgramData")), Path(os.environ.get("APPDATA", ""))]
    best = None
    for root in roots:
        for link in glob.glob(str(root / "Microsoft" / "Windows" / "Start Menu" / "Programs" / "**" / "*.lnk"), recursive=True):
            stem = Path(link).stem.lower()
            if stem == name:
                return Path(link)
            if name in stem and (best is None or len(stem) < len(best.stem)):
                best = Path(link)
    return best


KNOWN_FOLDERS = {"downloads": "Downloads", "download": "Downloads", "documents": "Documents", "docs": "Documents",
                 "desktop": "Desktop", "pictures": "Pictures", "photos": "Pictures", "music": "Music", "videos": "Videos",
                 "home": "", "user folder": ""}


def known_folder(name: str) -> Path | None:
    """'downloads', 'my pcs download folder', 'the documents folder' -> that folder in the home
    directory (or its OneDrive copy, where Windows has moved it). None for anything else."""
    words = re.sub(r"\b(my|the|pc'?s|pcs|computer'?s|laptop'?s|folder|directory)\b", " ", name.lower())
    key = " ".join(words.split())
    if key not in KNOWN_FOLDERS:
        return None
    home = Path.home()
    sub = KNOWN_FOLDERS[key]
    for candidate in (home / sub, home / "OneDrive" / sub):
        if candidate.is_dir():
            return candidate
    return None


def open(target: str, **_: object) -> dict:  # noqa: A001 - the action's name
    target = target.strip()
    domain = re.fullmatch(r"[\w-]+(?:\.[\w-]+)*\.([a-z]{2,})(?:/\S*)?", target.lower())
    bare_domain = bool(domain) and domain.group(1) not in FILE_EXTENSIONS and not _path(target).exists()
    if target.lower().startswith(("http://", "https://", "www.")) or bare_domain:
        url = target if "://" in target else "https://" + target
        webbrowser.open(url)
        return {"opened": url, "kind": "web page"}
    path = _path(target)
    if path.exists():
        os.startfile(str(path))  # noqa: S606 - the user asked to open it
        return {"opened": str(path), "kind": "folder" if path.is_dir() else "file"}
    name = target.lower().removesuffix(".exe")
    folder = known_folder(name)
    if folder is not None:  # "Downloads", "my documents folder": the folder, not an app called that
        os.startfile(str(folder))  # noqa: S606
        return {"opened": str(folder), "kind": "folder"}
    page = name.removesuffix(" settings").removeprefix("the ").strip()
    if name.endswith(" settings") and page.replace("-", "").replace(" ", "") in SETTINGS_PAGES:  # "sound settings"
        os.startfile(f"ms-settings:{SETTINGS_PAGES[page.replace('-', '').replace(' ', '')]}")  # noqa: S606
        return {"opened": f"Settings: {page}", "kind": "settings"}
    command = APPS.get(name)
    if command is None and name.removesuffix(" website").removesuffix(" site") in SITES:
        url = SITES[name.removesuffix(" website").removesuffix(" site")]
        webbrowser.open(url)
        return {"opened": url, "kind": "web page"}
    if command is None:
        shortcut = _start_menu_shortcut(name)
        if shortcut is None:
            raise ActionError(f"couldn't find an app called {target!r}; for a website, use pc_go_to")
        os.startfile(str(shortcut))  # noqa: S606
        return {"opened": shortcut.stem, "kind": "app"}
    try:
        os.startfile(command)  # noqa: S606 - a fixed app name or protocol from APPS
    except OSError:
        shortcut = _start_menu_shortcut(name)
        if shortcut is None:
            raise ActionError(f"{target} doesn't seem to be installed") from None
        os.startfile(str(shortcut))  # noqa: S606
    return {"opened": target, "kind": "app"}


def _speakers():  # noqa: ANN202 - pycaw's IAudioEndpointVolume
    """The default output's volume control (Windows Core Audio), or None without pycaw."""
    try:
        import comtypes
        from pycaw.pycaw import AudioUtilities
    except ImportError:
        return None
    comtypes.CoInitialize()  # each action runs on its own thread
    device = AudioUtilities.GetSpeakers()
    return getattr(device, "EndpointVolume", None)


def volume_level() -> tuple[int, bool] | None:
    """(volume percent, muted) now, or None when it can't be read."""
    speakers = _speakers()
    if speakers is None:
        return None
    return round(speakers.GetMasterVolumeLevelScalar() * 100), bool(speakers.GetMute())


def media(action: str, steps: int = 1, **_: object) -> dict:
    before = volume_level() if action in ("volume_up", "volume_down", "mute") else None
    _key(VK[action], steps if action in ("volume_up", "volume_down") else 1)
    if action in ("volume_up", "volume_down", "mute"):
        out: dict = {"done": action, "steps": steps}
        after = volume_level()
        if after is not None:
            out.update(volume=after[0], muted=after[1])
        if before is not None and action == "mute":
            out["undo"] = {"function": "media", "args": {"action": "mute"}, "label": "unmute" if not before[1] else "mute again"}
        elif before is not None:
            out["undo"] = {"function": "set_volume", "args": {"level": before[0]}, "label": f"put the volume back to {before[0]}%"}
        return out
    time.sleep(1.0)  # let the player update its title
    return {"done": action, **now_playing()}


def set_volume(level: int, **_: object) -> dict:
    speakers = _speakers()
    if speakers is None:  # no Core Audio: key presses, 2% each
        _key(VK["volume_down"], 50)
        _key(VK["volume_up"], round(level / 2))
        return {"volume": round(level / 2) * 2}
    before = round(speakers.GetMasterVolumeLevelScalar() * 100)
    speakers.SetMasterVolumeLevelScalar(level / 100, None)
    if level > 0 and speakers.GetMute():
        speakers.SetMute(0, None)  # "volume 40" means you want to hear it
    return {"volume": level, "undo": {"function": "set_volume", "args": {"level": before},
                                      "label": f"put the volume back to {before}%"}}


def lock_screen(**_: object) -> dict:
    ctypes.windll.user32.LockWorkStation()
    return {"locked": True}


def best_window(name: str) -> tuple[str, str, int, int] | None:
    """The window someone means by `name`: the app itself ("WhatsApp" is WhatsApp, not a file
    called whatsapp_bot.py in an editor), else the closest title."""
    wanted = name.strip().lower()
    scored = []
    for title, process, pid, hwnd in _windows():
        t, app = title.lower(), process.lower().removesuffix(".exe").split(".")[0]
        if t == wanted or app == wanted:
            score = 4
        elif t.endswith(f" - {wanted}") or t.endswith(f" - {wanted} - google chrome") or app.startswith(wanted):
            score = 3
        elif re.search(rf"(^|[\s\-|]){re.escape(wanted)}($|[\s\-|])", t):
            score = 2
        elif wanted in t:
            score = 1
        else:
            continue
        scored.append((score, title, process, pid, hwnd))
    if not scored:
        return None
    _score, title, process, pid, hwnd = max(scored, key=lambda s: s[0])
    return title, process, pid, hwnd


def _window(title: str):  # noqa: ANN202
    import pygetwindow

    best = best_window(title)
    if best is None:
        raise ActionError(f"no open window matches {title!r}")
    return next(w for w in pygetwindow.getAllWindows() if w._hWnd == best[3])  # noqa: SLF001


def focus_window(title: str, **_: object) -> dict:
    window = _window(title)
    if window.isMinimized:
        window.restore()
    window.activate()
    return {"focused": window.title}


def close_window(title: str, **_: object) -> dict:
    window = _window(title)
    window.close()
    return {"closed": window.title}


def type_text(text: str, **_: object) -> dict:
    import pyautogui
    import pyperclip

    if text.isascii():
        pyautogui.write(text, interval=0.01)
    else:  # pyautogui types ASCII only: paste anything else, then put the clipboard back
        before = pyperclip.paste()
        pyperclip.copy(text)
        pyautogui.hotkey("ctrl", "v")
        time.sleep(0.2)
        pyperclip.copy(before)
    return {"typed_characters": len(text)}


KEY_NAMES = {"win": "winleft", "windows": "winleft", "control": "ctrl", "escape": "esc", "return": "enter", "del": "delete"}


def press_keys(keys: str, **_: object) -> dict:
    import pyautogui

    parts = [KEY_NAMES.get(k.strip().lower(), k.strip().lower()) for k in keys.replace(" ", "").split("+") if k.strip()]
    invalid = [k for k in parts if k not in pyautogui.KEYBOARD_KEYS]
    if invalid:
        raise ActionError(f"unknown key(s): {', '.join(invalid)}")
    pyautogui.hotkey(*parts)
    return {"pressed": "+".join(parts)}


def click(x: int, y: int, button: str = "left", **_: object) -> dict:
    import pyautogui

    width, height = pyautogui.size()
    if x >= width or y >= height:
        raise ActionError(f"({x}, {y}) is off the screen ({width}x{height})")
    if button == "double":
        pyautogui.doubleClick(x, y)
    else:
        pyautogui.click(x, y, button=button)
    return {"clicked": [x, y], "button": button}


UNDO_DIR = Path(os.environ.get("APPDATA", Path.home())) / "Custodian" / "undo"  # copies of files before an overwrite


def write_file(path: str, content: str, overwrite: bool = False, *, protected: list[str], **_: object) -> dict:
    file = _path(path)
    _guard(file, protected)
    if file.exists() and not overwrite:
        raise ActionError(f"{file} already exists; ask to overwrite it")
    undo = {"function": "recycle_created", "args": {"path": str(file)}, "label": f"move the new file {file.name} to the Recycle Bin"}
    if file.is_file():  # keep what was there, so "undo" can put it back
        UNDO_DIR.mkdir(parents=True, exist_ok=True)
        backup = UNDO_DIR / f"{dt.datetime.now():%Y%m%d-%H%M%S}-{file.name}"
        shutil.copy2(file, backup)
        undo = {"function": "restore_backup", "args": {"backup": str(backup), "path": str(file)},
                "label": f"put back the old {file.name}"}
    file.parent.mkdir(parents=True, exist_ok=True)
    file.write_text(content, encoding="utf-8")
    return {"written": str(file), "characters": len(content), "undo": undo}


def move_file(source: str, destination: str, *, protected: list[str], **_: object) -> dict:
    src, dest = _path(source), _path(destination)
    _guard(src, protected)
    _guard(dest, protected)
    if not src.exists():
        raise ActionError(f"{src} doesn't exist")
    if dest.is_dir():
        dest = dest / src.name
    if dest.exists():
        raise ActionError(f"{dest} already exists")
    shutil.move(str(src), str(dest))
    return {"moved": str(src), "to": str(dest),
            "undo": {"function": "move_file", "args": {"source": str(dest), "destination": str(src)},
                     "label": f"move {dest.name} back to {src.parent}"}}


def delete_file(path: str, *, protected: list[str], **_: object) -> dict:
    from send2trash import send2trash

    target = _path(path)
    _guard(target, protected)
    if target == Path.home() or len(target.parts) <= 2:
        raise ActionError(f"Custodian won't delete {target}")
    if not target.exists():
        raise ActionError(f"{target} doesn't exist")
    send2trash(str(target))
    return {"moved_to_recycle_bin": str(target),
            "undo": {"function": "restore_deleted", "args": {"path": str(target)},
                     "label": f"restore {target.name} from the Recycle Bin"}}


# Undoing (only ever sent by the server, from a record an action above made) ----------------------

def _recycle_bin():  # noqa: ANN202 - a Shell folder
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    return win32com.client.Dispatch("Shell.Application").Namespace(10)  # ssfBITBUCKET


def restore_deleted(path: str, *, protected: list[str], **_: object) -> dict:
    """Put a file Custodian moved to the Recycle Bin back where it was."""
    target = _path(path)
    _guard(target, protected)
    if target.exists():
        raise ActionError(f"{target} is already there")
    folder = _recycle_bin()
    found = []
    for item in folder.Items():
        stored = Path(item.Path)  # C:\$Recycle.Bin\<sid>\$R1A2B3C.pdf: the file itself, with its extension
        same_place = os.path.normcase(folder.GetDetailsOf(item, 1)) == os.path.normcase(str(target.parent))  # "Original Location"
        same_name = item.Name.lower() in (target.name.lower(), target.stem.lower()) and stored.suffix.lower() == target.suffix.lower()
        if same_place and same_name:
            info = stored.with_name("$I" + stored.name[2:])  # its record, written when it was deleted
            found.append((info.stat().st_mtime if info.exists() else 0.0, item))
    if not found:
        raise ActionError(f"{target.name} isn't in the Recycle Bin any more")
    max(found, key=lambda pair: pair[0])[1].InvokeVerb("undelete")  # the newest deletion of that name
    for _ in range(30):
        if target.exists():
            return {"restored": str(target)}
        time.sleep(0.1)
    raise ActionError(f"Windows didn't restore {target.name}")


def undo_moves(moves: list, *, protected: list[str], **_: object) -> dict:
    """Put files back where they were before a sort: `moves` is [[now at, was at], ...]."""
    back, skipped, emptied = 0, [], set()
    for now_at, was_at in moves[:2000]:
        src, dest = _path(now_at), _path(was_at)
        _guard(src, protected)
        _guard(dest, protected)
        if not src.exists() or dest.exists():
            skipped.append(src.name)
            continue
        shutil.move(str(src), str(dest))
        back += 1
        emptied.add(src.parent)
    for folder in emptied:  # the subfolders the sort made, if nothing else is in them now
        try:
            folder.rmdir()
        except OSError:
            pass
    return {"moved_back": back, "skipped": skipped[:20] or None}


def remove_empty_folder(path: str, *, protected: list[str], **_: object) -> dict:
    folder = _path(path)
    _guard(folder, protected)
    if not folder.is_dir():
        raise ActionError(f"{folder} isn't there any more")
    if any(folder.iterdir()):
        raise ActionError(f"{folder.name} isn't empty any more, so Custodian left it")
    folder.rmdir()
    return {"removed": str(folder)}


def recycle_created(path: str, *, protected: list[str], **_: object) -> dict:
    """Move a file or folder Custodian created (a copy, a new file) to the Recycle Bin."""
    from send2trash import send2trash

    target = _path(path)
    _guard(target, protected)
    if not target.exists():
        raise ActionError(f"{target} isn't there any more")
    send2trash(str(target))
    return {"moved_to_recycle_bin": str(target)}


def restore_backup(backup: str, path: str, *, protected: list[str], **_: object) -> dict:
    saved, target = Path(backup).resolve(), _path(path)
    _guard(target, protected)
    if UNDO_DIR.resolve() not in saved.parents or not saved.is_file():
        raise ActionError("the saved copy is gone")
    shutil.copy2(saved, target)
    return {"restored": str(target)}


def set_clipboard(text: str, **_: object) -> dict:
    import pyperclip

    pyperclip.copy(text)
    return {"copied_characters": len(text)}


# The web -----------------------------------------------------------------------------------------

def go_to(query: str, **_: object) -> dict:
    """The first search result, opened directly (DuckDuckGo's "\\" jumps to it)."""
    from urllib.parse import quote

    known = SITES.get(query.strip().lower())
    url = known or f"https://duckduckgo.com/?q={quote(chr(92) + query.strip())}"
    webbrowser.open(url)
    return {"opened": url, "for": query, "kind": "web page"}


def search_site(site: str, query: str, **_: object) -> dict:
    from urllib.parse import quote, quote_plus

    template = SEARCH_URLS[site]
    url = template.format(quote(query) if site in ("spotify", "gmail", "maps") else quote_plus(query))
    webbrowser.open(url)
    return {"opened": url, "searched": site, "for": query}


# Music and video ---------------------------------------------------------------------------------

def _windows() -> list[tuple[str, str, int, int]]:
    """(title, process name, pid, hwnd) of every visible window with a title."""
    import psutil

    user32 = ctypes.windll.user32
    found: list[tuple[str, str, int, int]] = []

    @ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)
    def each(hwnd, _lparam):  # noqa: ANN001, ANN202
        if user32.IsWindowVisible(hwnd):
            length = user32.GetWindowTextLengthW(hwnd)
            if length:
                buffer = ctypes.create_unicode_buffer(length + 1)
                user32.GetWindowTextW(hwnd, buffer, length + 1)
                pid = ctypes.c_ulong()
                user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
                try:
                    name = psutil.Process(pid.value).name()
                except (psutil.Error, ValueError):
                    name = "?"
                found.append((buffer.value, name, pid.value, hwnd))
        return True

    user32.EnumWindows(each, 0)
    return found


def now_playing(**_: object) -> dict:
    playing = []
    for title, process, _pid, _hwnd in _windows():
        if process.lower() == "spotify.exe" and title not in ("Spotify", "Spotify Premium", "Spotify Free"):
            playing.append({"app": "Spotify", "title": title})
        elif " - YouTube" in title:
            playing.append({"app": "YouTube", "title": title.split(" - YouTube")[0]})
    spotify_paused = any(p.lower() == "spotify.exe" and t in ("Spotify Premium", "Spotify Free") for t, p, _, _ in _windows())
    return {"now_playing": playing or None, **({"spotify": "paused or idle"} if spotify_paused else {})}


def play(query: str, service: str = "youtube", **_: object) -> dict:
    from urllib.parse import quote, quote_plus

    if service == "spotify":
        os.startfile(f"spotify:search:{quote(query)}")  # noqa: S606 - Spotify's own link scheme
        started = _spotify_play_top_result()
        time.sleep(1.5)
        return {"service": "spotify", "searched": query, "started": started, **now_playing()}
    import httpx

    try:
        page = httpx.get(f"https://www.youtube.com/results?search_query={quote_plus(query)}", timeout=10,
                         headers={"User-Agent": "Mozilla/5.0", "Accept-Language": "en-US,en;q=0.9"}).text
    except httpx.HTTPError as exc:
        raise ActionError(f"couldn't reach YouTube ({exc})") from None
    match = re.search(r'"videoRenderer":\{"videoId":"([\w-]{11})".*?"title":\{"runs":\[\{"text":"((?:[^"\\]|\\.)*)"', page)
    if not match:
        url = f"https://www.youtube.com/results?search_query={quote_plus(query)}"
        webbrowser.open(url)
        return {"service": "youtube", "opened": url, "started": False, "note": "no video found; showing the search results"}
    video, title = match.group(1), json_unescape(match.group(2))
    url = f"https://www.youtube.com/watch?v={video}"
    webbrowser.open(url)
    return {"service": "youtube", "playing": title, "url": url, "started": True}


def json_unescape(text: str) -> str:
    import json

    try:
        return json.loads(f'"{text}"')
    except ValueError:
        return text


def _spotify_play_top_result() -> bool:
    """Spotify opens its search; press the top result's Play button."""
    try:
        from pywinauto import Desktop
    except ImportError:
        return False
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        time.sleep(1.0)
        try:
            window = next((w for w in Desktop(backend="uia").windows() if w.process_id() in
                           {pid for _t, p, pid, _h in _windows() if p.lower() == "spotify.exe"}), None)
            if window is None:
                continue
            buttons = [b for b in window.descendants(control_type="Button")
                       if b.window_text().lower().startswith("play") and b.window_text().lower() != "play"]
            if buttons:
                buttons[0].invoke()
                return True
        except Exception:  # noqa: BLE001 - Spotify still loading
            continue
    return False


# Settings ----------------------------------------------------------------------------------------

def set_brightness(level: int, **_: object) -> dict:
    try:
        import screen_brightness_control as sbc

        before = sbc.get_brightness()
        sbc.set_brightness(level)
    except Exception as exc:  # noqa: BLE001 - e.g. an external monitor without DDC/CI
        raise ActionError(f"this screen's brightness can't be set from Windows ({exc})") from None
    out: dict = {"brightness": level}
    if before:
        out["undo"] = {"function": "set_brightness", "args": {"level": int(before[0])},
                       "label": f"put the brightness back to {int(before[0])}%"}
    return out


def set_theme(mode: str, **_: object) -> dict:
    import winreg

    path = r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize"
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, path) as key:
            before = "light" if winreg.QueryValueEx(key, "AppsUseLightTheme")[0] else "dark"
    except OSError:
        before = None
    with winreg.OpenKey(winreg.HKEY_CURRENT_USER, path, 0, winreg.KEY_SET_VALUE) as key:
        for name in ("AppsUseLightTheme", "SystemUsesLightTheme"):
            winreg.SetValueEx(key, name, 0, winreg.REG_DWORD, 0 if mode == "dark" else 1)
    out: dict = {"theme": mode}
    if before and before != mode:
        out["undo"] = {"function": "set_theme", "args": {"mode": before}, "label": f"switch back to {before} mode"}
    return out


def open_settings(page: str, **_: object) -> dict:
    os.startfile(f"ms-settings:{SETTINGS_PAGES[page]}")  # noqa: S606 - a fixed Settings page
    return {"opened": f"Settings: {page}"}


_clocks: dict[str, object] = {}  # timer or alarm id -> threading.Timer
ALARMS_FILE = Path(os.environ.get("APPDATA", Path.home())) / "Custodian" / "alarms.json"


def _ring(label: str) -> None:
    if on_timer is not None:
        on_timer(label)
        return
    try:
        import winsound

        for _ in range(3):
            winsound.Beep(988, 300)
    except (ImportError, RuntimeError):
        pass


def _start_clock(clock_id: str, seconds: float, label: str, *, alarm: bool = False) -> None:
    import threading

    def fire() -> None:
        _clocks.pop(clock_id, None)
        if alarm:
            _save_alarms([a for a in _load_alarms() if a["id"] != clock_id])
        _ring(label)

    clock = threading.Timer(max(0.0, seconds), fire)
    clock.daemon = True
    clock.start()
    _clocks[clock_id] = clock


def timer(seconds: int, label: str = "Timer", **_: object) -> dict:
    clock_id = f"timer-{time.time_ns()}"
    _start_clock(clock_id, seconds, label)
    ends = dt.datetime.now() + dt.timedelta(seconds=seconds)
    return {"timer": label, "rings_at": ends.strftime("%H:%M:%S"), "seconds": seconds,
            "undo": {"function": "cancel_timer", "args": {"id": clock_id}, "label": f"cancel the {label.lower()}"}}


def _load_alarms() -> list[dict]:
    import json

    try:
        return json.loads(ALARMS_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []


def _save_alarms(alarms: list[dict]) -> None:
    import json

    ALARMS_FILE.parent.mkdir(parents=True, exist_ok=True)
    ALARMS_FILE.write_text(json.dumps(alarms), encoding="utf-8")


def alarm(time: str, label: str = "Alarm", **_: object) -> dict:  # noqa: A002 - the action's argument name
    """An alarm at the next `time` (HH:MM) - today if it's still to come, else tomorrow. Kept in
    alarms.json, so it rings even if the app restarts before then (restore_alarms)."""
    hour, minute = (int(part) for part in time.split(":"))
    now = dt.datetime.now()
    when = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if when <= now:
        when += dt.timedelta(days=1)
    clock_id = f"alarm-{when:%Y%m%d%H%M}-{abs(hash(label)) % 10000}"
    _save_alarms([a for a in _load_alarms() if a["id"] != clock_id] + [{"id": clock_id, "at": when.isoformat(), "label": label}])
    _start_clock(clock_id, (when - now).total_seconds(), label, alarm=True)
    day = "today" if when.date() == now.date() else "tomorrow"
    return {"alarm": label, "label": label, "rings_at": f"{when:%H:%M} {day}", "id": clock_id,
            "undo": {"function": "cancel_timer", "args": {"id": clock_id}, "label": f"cancel the {when:%H:%M} alarm"}}


def restore_alarms() -> int:
    """At start-up: set again every saved alarm still to come (one missed by less than 10 minutes rings now)."""
    now, kept = dt.datetime.now(), []
    for saved in _load_alarms():
        when = dt.datetime.fromisoformat(saved["at"])
        if when < now - dt.timedelta(minutes=10):
            continue
        kept.append(saved)
        _start_clock(saved["id"], (when - now).total_seconds(), saved["label"], alarm=True)
    _save_alarms(kept)
    return len(kept)


def cancel_timer(id: str, **_: object) -> dict:  # noqa: A002
    clock = _clocks.pop(id, None)
    _save_alarms([a for a in _load_alarms() if a["id"] != id])
    if clock is None:
        raise ActionError("that timer or alarm already rang or was cancelled")
    clock.cancel()
    return {"cancelled": id}


# Do Not Disturb ----------------------------------------------------------------------------------
# Windows 11 has no supported API for it, so this uses its own switch the way a screen reader
# would: the "Do not disturb" button in the notification centre (Win+N), read and toggled through
# UI Automation, then the centre is closed again.

def _dnd_toggle():  # noqa: ANN202 - (UIA element, its TogglePattern, the notification centre's handle)
    from pywinauto.uia_defines import IUIA

    iuia, ids = IUIA().iuia, IUIA().UIA_dll
    root = iuia.GetRootElement()
    centre = iuia.CreatePropertyConditionEx(ids.UIA_NamePropertyId, "Notification", 3)  # "Notification Center", any case
    condition = iuia.CreateAndCondition(iuia.CreatePropertyCondition(ids.UIA_NamePropertyId, "Do not disturb"),
                                        iuia.CreatePropertyCondition(ids.UIA_IsTogglePatternAvailablePropertyId, True))
    for _ in range(25):
        window = root.FindFirst(ids.TreeScope_Children, centre)  # the centre's own window: a small tree to search
        found = window.FindFirst(ids.TreeScope_Descendants, condition) if window else None
        if found:
            pattern = found.GetCurrentPattern(ids.UIA_TogglePatternId).QueryInterface(ids.IUIAutomationTogglePattern)
            return found, pattern
        time.sleep(0.1)
    return None, None


def dnd_state() -> bool | None:
    """Whether Do Not Disturb is on, or None if it couldn't be read."""
    _keystroke(VK_LWIN, VK_N)
    try:
        _element, pattern = _dnd_toggle()
        return None if pattern is None else pattern.CurrentToggleState == 1
    finally:
        _keystroke(VK_ESCAPE)


def do_not_disturb(on: bool = True, **_: object) -> dict:
    _keystroke(VK_LWIN, VK_N)
    try:
        _element, pattern = _dnd_toggle()
        if pattern is None:
            raise ActionError("couldn't find Windows' Do not disturb switch; it's in the notification centre (Win+N)")
        before = pattern.CurrentToggleState == 1
        if before != on:
            pattern.Toggle()
            time.sleep(0.3)
        now = pattern.CurrentToggleState == 1
    finally:
        _keystroke(VK_ESCAPE)
    if now != on:
        raise ActionError("Windows didn't change Do not disturb")
    out: dict = {"do_not_disturb": now}
    if before != now:
        out["undo"] = {"function": "do_not_disturb", "args": {"on": before},
                       "label": "turn Do Not Disturb " + ("on" if before else "off") + " again"}
    return out


# Speaking and routines ---------------------------------------------------------------------------

on_say = None  # set by the app: speaks a sentence with the reply voice


def say(text: str, **_: object) -> dict:
    if on_say is None:
        raise ActionError("this desktop app can't speak")
    on_say(text)
    return {"said": text}


def confirm_plan(title: str, steps: list[str], **_: object) -> dict:
    """Nothing to do by itself: the dialog that asked about the whole list was the point."""
    return {"approved": True, "steps": len(steps)}


# Apps --------------------------------------------------------------------------------------------

def list_apps(**_: object) -> dict:
    import psutil

    apps: dict[str, dict] = {}
    for title, process, pid, _hwnd in _windows():
        entry = apps.setdefault(process, {"app": process.removesuffix(".exe"), "windows": 0, "memory_mb": 0})
        entry["windows"] += 1
        try:
            entry["memory_mb"] = max(entry["memory_mb"], round(psutil.Process(pid).memory_info().rss / 1e6))
        except psutil.Error:
            pass
    return {"apps": sorted(apps.values(), key=lambda a: -a["memory_mb"])[:40]}


APP_ALIASES = {"word": "winword", "excel": "excel", "powerpoint": "powerpnt", "vs code": "code", "vscode": "code",
               "edge": "msedge", "file explorer": "explorer", "whatsapp": "whatsapp", "teams": "ms-teams"}


def app_process(name: str) -> str | None:
    """The running process someone means by an app name: exact first, then the shortest
    process name starting with it, then one containing it."""
    wanted = APP_ALIASES.get(name.strip().lower(), name.strip().lower().removesuffix(".exe"))
    names = {p for _t, p, _pid, _h in _windows()}
    stems = {p: p.lower().removesuffix(".exe") for p in names}
    for test in (lambda s: s == wanted or s.split(".")[0] == wanted, lambda s: s.startswith(wanted), lambda s: wanted in s):
        matches = sorted((p for p, s in stems.items() if test(s)), key=len)
        if matches:
            return matches[0]
    return None


def foreground_process() -> str:
    """The process name of the app in front ("chrome.exe"), or "" when it can't be told."""
    import psutil

    pid = ctypes.c_ulong()
    ctypes.windll.user32.GetWindowThreadProcessId(ctypes.windll.user32.GetForegroundWindow(), ctypes.byref(pid))
    try:
        return psutil.Process(pid.value).name().lower()
    except (psutil.Error, ValueError):
        return ""


def memory_hogs(front_seen: dict[str, float], now: float, *, idle_minutes: float = 10, min_mb: int = 300,
                limit: int = 2, windows: list | None = None, processes: list | None = None) -> list[dict]:
    """The apps worth closing when memory runs out: the ones with a window that hold the most
    memory (all their processes together - Chrome is dozens) and haven't been in front for
    `idle_minutes` (`front_seen`: process name -> when it last was). Never Windows' own processes,
    this app, or the app in front."""
    import psutil

    windows = _windows() if windows is None else windows
    with_window = {p.lower() for _t, p, _pid, _h in windows} - SYSTEM_PROCESSES - {"?"}
    candidates = {p for p in with_window if now - front_seen.get(p, -1e9) >= idle_minutes * 60}
    held: dict[str, int] = {}
    for proc in psutil.process_iter(["name", "memory_info"]) if processes is None else processes:
        name = (proc.info.get("name") or "").lower()
        if name in candidates and proc.info.get("memory_info"):
            held[name] = held.get(name, 0) + proc.info["memory_info"].rss
    hogs = [{"process": n, "app": n.removesuffix(".exe"), "memory_mb": round(b / 1e6)} for n, b in held.items()
            if b / 1e6 >= min_mb]
    return sorted(hogs, key=lambda a: -a["memory_mb"])[:limit]


def close_app(name: str, **_: object) -> dict:
    import psutil

    process = app_process(name)
    targets = [(t, p, pid, h) for t, p, pid, h in _windows() if p == process]
    if not targets:
        raise ActionError(f"no running app with a window matches {name!r}")
    if any(p.lower() in SYSTEM_PROCESSES for _t, p, _pid, _h in targets):
        raise ActionError(f"Custodian won't close {targets[0][1]}: Windows or Custodian itself needs it")
    for _title, _process, _pid, hwnd in targets:
        ctypes.windll.user32.PostMessageW(hwnd, 0x0010, 0, 0)  # WM_CLOSE: the app can still ask to save
    time.sleep(3)
    still = [t for t, p, pid, _h in _windows() if pid in {pid for _t, _p, pid, _h in targets}]
    pids = {pid for _t, _p, pid, _h in targets}
    out = {"closed": targets[0][1].removesuffix(".exe"), "windows_closed": len(targets) - len(still),
           "still_open": still or None, "running": any(psutil.pid_exists(p) for p in pids) if still else False}
    if len(still) < len(targets):
        out["undo"] = {"function": "open", "args": {"target": name}, "label": f"open {name} again"}
    return out


# Files -------------------------------------------------------------------------------------------

def copy_file(source: str, destination: str, *, protected: list[str], **_: object) -> dict:
    src, dest = _path(source), _path(destination)
    _guard(dest, protected)
    if not src.exists():
        raise ActionError(f"{src} doesn't exist")
    if dest.is_dir():
        dest = dest / src.name
    if dest.exists():
        raise ActionError(f"{dest} already exists")
    if src.is_dir():
        shutil.copytree(src, dest)
    else:
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dest)
    return {"copied": str(src), "to": str(dest),
            "undo": {"function": "recycle_created", "args": {"path": str(dest)}, "label": f"move the copy {dest.name} to the Recycle Bin"}}


def create_folder(path: str, *, protected: list[str], **_: object) -> dict:
    folder = _path(path)
    _guard(folder, protected)
    existed = folder.exists()
    folder.mkdir(parents=True, exist_ok=True)
    out: dict = {"folder": str(folder), "already_existed": existed}
    if not existed:
        out["undo"] = {"function": "remove_empty_folder", "args": {"path": str(folder)}, "label": f"remove the new folder {folder.name}"}
    return out


def organize_folder(path: str, *, protected: list[str], **_: object) -> dict:
    folder = _path(path)
    _guard(folder, protected)
    if not folder.is_dir():
        raise ActionError(f"{folder} isn't a folder")
    moved: dict[str, int] = {}
    moves: list[list[str]] = []  # [now at, was at], for undo
    for entry in folder.iterdir():
        if not entry.is_file() or entry.name.startswith(".") or entry.suffix.lower() in (".crdownload", ".part", ".tmp", ".lnk", ".ini"):
            continue
        ext = entry.suffix.lower().lstrip(".")
        kind = next((k for k, exts in FILE_TYPES.items() if ext in exts), "Other")
        target_dir = folder / kind
        target_dir.mkdir(exist_ok=True)
        target, n = target_dir / entry.name, 2
        while target.exists():
            target, n = target_dir / f"{entry.stem} ({n}){entry.suffix}", n + 1
        shutil.move(str(entry), str(target))
        moved[kind] = moved.get(kind, 0) + 1
        moves.append([str(target), str(entry)])
    out: dict = {"folder": str(folder), "moved": moved, "files_moved": sum(moved.values())}
    if moves and len(moves) <= 2000:
        out["undo"] = {"function": "undo_moves", "args": {"moves": moves},
                       "label": f"put the {len(moves)} sorted files back in {folder.name}"}
    return out


# Writing into apps -------------------------------------------------------------------------------

def write_in_app(text: str, app: str = "notepad", **_: object) -> dict:
    import subprocess

    import pyautogui
    import pyperclip

    before = {h for _t, _p, _pid, h in _windows()}
    command = {"notepad": ["notepad.exe"], "wordpad": ["write.exe"], "word": ["cmd", "/c", "start", "", "winword", "/n"]}[app]
    subprocess.Popen(command)  # noqa: S603 - a fixed app
    marker = {"notepad": "notepad", "wordpad": "wordpad", "word": "word"}[app]
    deadline, window = time.monotonic() + 15, None
    while time.monotonic() < deadline and window is None:
        time.sleep(0.5)
        window = next(((t, h) for t, p, _pid, h in _windows() if h not in before and marker in (t + p).lower()), None)
    if window is None:
        raise ActionError(f"{app} didn't open a new window")
    ctypes.windll.user32.SetForegroundWindow(window[1])
    time.sleep(0.8 if app == "notepad" else 2.5)
    saved = pyperclip.paste()
    pyperclip.copy(text)
    pyautogui.hotkey("ctrl", "v")
    time.sleep(0.3)
    pyperclip.copy(saved)
    return {"app": app, "window": window[0], "characters": len(text)}


def _link(paragraph, url: str, text: str) -> None:  # noqa: ANN001 - a python-docx paragraph
    """A clickable link in a Word paragraph (python-docx has no call for one)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                   is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    for tag, attrs in (("w:color", {"w:val": "0563C1"}), ("w:u", {"w:val": "single"})):
        element = OxmlElement(tag)
        for key, value in attrs.items():
            element.set(qn(key), value)
        props.append(element)
    run.append(props)
    label = OxmlElement("w:t")
    label.text = text
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def sections_of(body: str) -> list[dict]:
    """'## Heading' (or '#', '###') lines start a section; text before the first has no heading."""
    out: list[dict] = [{"heading": "", "text": ""}]
    for line in body.splitlines():
        heading = re.match(r"^\s{0,3}#{1,4}\s+(.+?)\s*#*\s*$", line)
        if heading:
            out.append({"heading": heading.group(1).strip("* "), "text": ""})
        else:
            out[-1]["text"] += line + "\n"
    return [s for s in out if s["heading"] or s["text"].strip()]


def write_document(title: str, sections: list[dict] | None = None, sources: list[dict] | None = None, body: str = "", *,
                   protected: list[str], **_: object) -> dict:
    """A Word document with the title, each section's heading and paragraphs ('- ' lines become
    bullets), and the sources as links - saved as a new file in Documents, then opened. The text
    comes as `body` ('## Heading' lines) or as `sections`."""
    from docx import Document

    sections = (sections or []) + (sections_of(body) if body.strip() else [])
    same = lambda a, b: re.sub(r"\W+", "", a).lower() == re.sub(r"\W+", "", b).lower()  # noqa: E731
    if sections and same(sections[0]["heading"], title):
        sections[0] = {**sections[0], "heading": ""}  # the title again as the first heading: once is enough
    if sources:  # the linked list below replaces a sources list written into the text
        sections = [s for s in sections if not re.fullmatch(r"(?i)\s*(sources?|references?)\s*:?\s*", s["heading"])]
    doc = Document()
    doc.add_heading(title, 0)
    for section in sections:
        if section["heading"]:
            doc.add_heading(section["heading"], 1)
        for block in re.split(r"\n\s*\n", section["text"].strip()):
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            if lines and all(line.startswith(("- ", "* ", "• ")) for line in lines):
                for line in lines:
                    doc.add_paragraph(line[2:].strip().replace("**", ""), style="List Bullet")
            elif lines:
                doc.add_paragraph(" ".join(lines).replace("**", ""))
    if sources:
        doc.add_heading("Sources", 1)
        for number, source in enumerate(sources, 1):
            paragraph = doc.add_paragraph(f"[{number}] ")
            _link(paragraph, source["url"], source["title"])
            paragraph.add_run(f" - {source['url']}")
    folder = known_folder("documents") or Path.home() / "Documents"
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", title).strip(" .")[:80] or "Custodian document"
    path, n = folder / f"{stem}.docx", 2
    while path.exists():  # never over another file
        path, n = folder / f"{stem} ({n}).docx", n + 1
    _guard(path, protected)
    folder.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    os.startfile(str(path))  # noqa: S606 - the file just made, in Word (or what opens .docx here)
    return {"saved": str(path), "opened": True, "sections": len(sections), "sources": len(sources or []),
            "undo": {"function": "recycle_created", "args": {"path": str(path)}, "label": f"delete {path.name}"}}


# Seeing and using the screen (Windows UI Automation: names, not pixels) --------------------------

ELEMENT_TYPES = {"Button", "Hyperlink", "Edit", "ComboBox", "CheckBox", "RadioButton", "MenuItem", "TabItem", "ListItem",
                 "SplitButton", "Document"}


def _target_window(window: str = ""):  # noqa: ANN202
    from pywinauto import Desktop

    if window:
        best = best_window(window)
        if best is None:
            raise ActionError(f"no open window matches {window!r}")
        hwnd = best[3]
    else:
        hwnd = ctypes.windll.user32.GetForegroundWindow()
    return Desktop(backend="uia").window(handle=hwnd).wrapper_object()


def _elements(win, limit: int = 400) -> list:  # noqa: ANN001
    found, deadline = [], time.monotonic() + 8
    for element in win.descendants():
        if time.monotonic() > deadline or len(found) >= limit:
            break
        try:
            kind, name = element.element_info.control_type, element.window_text().strip()
        except Exception:  # noqa: BLE001
            continue
        if name and (kind in ELEMENT_TYPES or kind == "Text"):
            found.append((element, kind, name))
    return found


BROWSERS = {"chrome.exe", "msedge.exe", "brave.exe", "opera.exe"}


def read_screen(window: str = "", **_: object) -> dict:
    win = _target_window(window)
    out = _read(win)
    process = next((p for _t, p, _pid, h in _windows() if h == win.handle), "").lower()
    if process in BROWSERS and not out["text"]:
        out["note"] = ("the browser isn't showing this page's contents to Windows, so only its toolbar is listed. "
                       "Run desktop\\run-desktop.ps1 -BrowserAccess once and restart the browser to allow it.")
    return out


def _read(win) -> dict:  # noqa: ANN001
    controls, texts, seen = [], [], set()
    for _element, kind, name in _elements(win):
        if (kind, name) in seen:
            continue
        seen.add((kind, name))
        if kind == "Text":
            if len(texts) < 40:
                texts.append(name[:200])
        elif len(controls) < 80:
            controls.append({"name": name[:120], "type": kind})
    return {"window": win.window_text(), "controls": controls, "text": texts}


def _find_element(name: str, window: str, types: set[str] | None = None):  # noqa: ANN202
    win = _target_window(window)
    candidates = [(e, k, n) for e, k, n in _elements(win) if k != "Text" and (types is None or k in types)]
    wanted = name.strip().lower()
    for test in (lambda n: n.lower() == wanted, lambda n: n.lower().startswith(wanted), lambda n: wanted in n.lower()):
        match = next((c for c in candidates if test(c[2])), None)
        if match:
            return win, match
    raise ActionError(f"nothing named {name!r} in {win.window_text()!r}; call pc_read_screen to see the names")


def click_element(name: str, window: str = "", **_: object) -> dict:
    win, (element, kind, found) = _find_element(name, window)
    try:
        element.invoke()
    except Exception:  # noqa: BLE001 - not every control supports Invoke
        element.click_input()
    return {"clicked": found, "type": kind, "window": win.window_text()}


def fill_field(name: str, text: str, window: str = "", **_: object) -> dict:
    win, (element, kind, found) = _find_element(name, window, {"Edit", "ComboBox", "Document"})
    try:
        element.set_edit_text(text)
    except Exception:  # noqa: BLE001 - e.g. a web page field: click it and paste
        import pyautogui
        import pyperclip

        element.click_input()
        saved = pyperclip.paste()
        pyperclip.copy(text)
        pyautogui.hotkey("ctrl", "a")
        pyautogui.hotkey("ctrl", "v")
        time.sleep(0.2)
        pyperclip.copy(saved)
    return {"filled": found, "type": kind, "characters": len(text), "window": win.window_text()}


# Messaging ---------------------------------------------------------------------------------------

confirm_hook = None  # set by the app: (title, details) -> bool, for actions that confirm what they found
COMPOSER_PREFIX = "type a message to "
# UI Automation's "that element isn't there any more": the app redrew it between finding it and using it, as
# WhatsApp does while its search results come in. (Windows' own text for the code, "An event was unable to
# invoke any of the subscribers", says nothing useful.)
UIA_E_ELEMENTNOTAVAILABLE = 0x80040201 - (1 << 32)


def _gone(exc: BaseException) -> bool:
    return type(exc).__name__ == "COMError" and bool(exc.args) and exc.args[0] == UIA_E_ELEMENTNOTAVAILABLE


class _Uia:
    """Native UI Automation lookups. WhatsApp (a web view) lists each element dozens of times,
    so walking its tree takes seconds; FindFirst with a condition takes milliseconds."""

    def __init__(self, hwnd: int) -> None:
        from pywinauto.uia_defines import IUIA

        self.iuia = IUIA().iuia
        self.ids = IUIA().UIA_dll
        self.root = self.iuia.ElementFromHandle(hwnd)

    def first(self, name: str | None = None, control: str | None = None, *, contains: bool = False):  # noqa: ANN201
        """The first visible element of `control` type whose name is (or contains) `name`."""
        ids, conditions = self.ids, [self.iuia.CreatePropertyCondition(self.ids.UIA_IsOffscreenPropertyId, False)]
        if name is not None:
            if contains:  # Windows 10 1809+: case-insensitive substring match
                conditions.append(self.iuia.CreatePropertyConditionEx(ids.UIA_NamePropertyId, name, 3))
            else:
                conditions.append(self.iuia.CreatePropertyCondition(ids.UIA_NamePropertyId, name))
        if control is not None:
            conditions.append(self.iuia.CreatePropertyCondition(ids.UIA_ControlTypePropertyId,
                                                                getattr(ids, f"UIA_{control}ControlTypeId")))
        condition = conditions[0]
        for extra in conditions[1:]:
            condition = self.iuia.CreateAndCondition(condition, extra)
        found = self.root.FindFirst(ids.TreeScope_Descendants, condition)
        return found if found else None  # "not found" comes back as a NULL pointer, not None

    def set_value(self, element, text: str) -> None:  # noqa: ANN001
        pattern = element.GetCurrentPattern(self.ids.UIA_ValuePatternId).QueryInterface(self.ids.IUIAutomationValuePattern)
        pattern.SetValue(text)

    def rows(self, limit: int = 40) -> list:
        """The visible list rows - WhatsApp's chats, or what its search found."""
        ids, found = self.ids, []
        for control in ("DataItem", "ListItem"):
            condition = self.iuia.CreateAndCondition(
                self.iuia.CreatePropertyCondition(ids.UIA_IsOffscreenPropertyId, False),
                self.iuia.CreatePropertyCondition(ids.UIA_ControlTypePropertyId, getattr(ids, f"UIA_{control}ControlTypeId")))
            items = self.root.FindAll(ids.TreeScope_Descendants, condition)
            found += [items.GetElement(i) for i in range(min(items.Length, limit))] if items else []
        return found[:limit]

    def press(self, element) -> None:  # noqa: ANN001
        """Activate a list row or button the way a screen reader does (no mouse involved)."""
        for pattern_id, interface, method in (
            (self.ids.UIA_InvokePatternId, self.ids.IUIAutomationInvokePattern, "Invoke"),
            (self.ids.UIA_SelectionItemPatternId, self.ids.IUIAutomationSelectionItemPattern, "Select"),
            (self.ids.UIA_LegacyIAccessiblePatternId, self.ids.IUIAutomationLegacyIAccessiblePattern, "DoDefaultAction"),
        ):
            try:
                pattern = element.GetCurrentPattern(pattern_id)
                if pattern:
                    getattr(pattern.QueryInterface(interface), method)()
                    return
            except Exception:  # noqa: BLE001 - try the next way
                continue
        from pywinauto.controls.uiawrapper import UIAWrapper
        from pywinauto.uia_element_info import UIAElementInfo

        UIAWrapper(UIAElementInfo(element)).click_input()

    def names(self, control: str, limit: int = 80) -> list[str]:
        """The names of the visible elements of a control type."""
        condition = self.iuia.CreateAndCondition(
            self.iuia.CreatePropertyCondition(self.ids.UIA_IsOffscreenPropertyId, False),
            self.iuia.CreatePropertyCondition(self.ids.UIA_ControlTypePropertyId, getattr(self.ids, f"UIA_{control}ControlTypeId")))
        items = self.root.FindAll(self.ids.TreeScope_Descendants, condition)
        found = []
        for i in range(min(items.Length, limit) if items else 0):
            try:
                found.append(items.GetElement(i).CurrentName or "")
            except Exception:  # noqa: BLE001 - redrawn meanwhile
                continue
        return found

    def expand(self, element) -> bool:  # noqa: ANN001
        """Open a menu button's menu - only that: Expand never presses the button. False when it has no menu."""
        try:
            pattern = element.GetCurrentPattern(self.ids.UIA_ExpandCollapsePatternId)
            if not pattern:
                return False
            pattern.QueryInterface(self.ids.IUIAutomationExpandCollapsePattern).Expand()
            return True
        except Exception:  # noqa: BLE001 - not a menu button, or gone
            return False


def _keystroke(*keys: int) -> None:
    """Press a key combination with keyboard events only: no mouse involved, so it works
    wherever the pointer is (pyautogui refuses when it's parked in a corner)."""
    user32 = ctypes.windll.user32
    for key in keys:
        user32.keybd_event(key, 0, 0, 0)
    for key in reversed(keys):
        user32.keybd_event(key, 0, 2, 0)
    time.sleep(0.05)


VK_CONTROL, VK_RETURN, VK_V, VK_A, VK_ESCAPE, VK_LWIN, VK_N = 0x11, 0x0D, 0x56, 0x41, 0x1B, 0x5B, 0x4E


def _paste_into_focus(text: str) -> None:
    import pyperclip

    saved = pyperclip.paste()
    pyperclip.copy(text)
    _keystroke(VK_CONTROL, VK_A)
    _keystroke(VK_CONTROL, VK_V)
    time.sleep(0.3)
    pyperclip.copy(saved)


_ROW_TAIL = re.compile(r"\s+(?:\d{1,2}:\d{2}\s?[ap]\.?m\.?|\d{1,2}:\d{2}|yesterday|today|\d{1,2}/\d{1,2}/\d{2,4}"
                       r"|monday|tuesday|wednesday|thursday|friday|saturday|sunday)\b.*$", re.I | re.S)


def _row_title(name: str) -> str:
    """'3 unread messages Tanishka 1:46 pm projects.csv' -> 'Tanishka'."""
    name = re.sub(r"^\d+ unread messages?\s+", "", name.strip(), flags=re.I)
    return _ROW_TAIL.sub("", name).strip()


def _same_person(wanted: str, chat: str) -> bool:
    """'tanishka' matches 'Tanishka' and 'Tanishka Sharma'; 'mum' doesn't match 'Mumbai Trip'."""
    want, have = re.findall(r"\w+", wanted.lower()), re.findall(r"\w+", chat.lower())
    return bool(want) and all(w in have for w in want)


NAME_MATCH = 0.85  # how well a chat's name must fit the name said to be picked without asking which one
NAME_CLOSE = 0.6  # ...and to be named as one they might mean ("Abhishek" for "Abhi", not "Zeenat" for "Zebra")
SEARCH_SECONDS = 3.0  # how long WhatsApp's search gets to list what it found
_NAME_FILLER = frozenset("on in to the my chat with whatsapp".split())


def _name_key(word: str) -> str:
    """How a name sounds, roughly, however it's spelled: 'Choudhury', 'Chaudhari' and 'Choudhary' are all
    'cdr'; 'Mronal' (how speech recognition wrote Mrunal) and 'Mrunal' are both 'mrnl'."""
    word = re.sub(r"[^a-z]", "", word.lower())
    for spelled, said in (("ph", "f"), ("w", "v"), ("z", "j"), ("ck", "k"), ("q", "k"), ("x", "ks")):
        word = word.replace(spelled, said)
    word = re.sub(r"(?<=[bcdgjkpstv])h", "", word)  # dh, bh, kh, th, ch: d, b, k, t, c
    if not word:
        return ""
    return re.sub(r"(.)\1+", r"\1", word[0] + re.sub(r"[aeiouy]", "", word[1:]))


def name_score(wanted: str, chat: str) -> float:
    """0-1: how well a chat's name fits a name as it was said - each word said must be there, spelled the same
    or nearly ('Abhi Choudhury' fits 'Abhi Chaudhari' 1.0, 'mum' fits 'Mumbai Trip' 0.67)."""
    import difflib

    said = [w for w in re.findall(r"\w+", wanted.lower()) if w not in _NAME_FILLER]
    have = re.findall(r"\w+", chat.lower())
    if not said or not have:
        return 0.0

    def fits(word: str) -> float:
        key = _name_key(word)
        return max(1.0 if len(key) >= 2 and key == _name_key(h) else difflib.SequenceMatcher(None, word, h).ratio()
                   for h in have)
    return sum(fits(w) for w in said) / len(said)


def _whatsapp_window() -> int:
    for attempt in range(30):
        found = next(((t, p, pid, h) for t, p, pid, h in _windows() if p.lower().startswith("whatsapp")), None)
        if found:
            return found[3]
        if attempt == 0:
            try:
                os.startfile("whatsapp:")  # noqa: S606 - WhatsApp's own link scheme
            except OSError:
                shortcut = _start_menu_shortcut("whatsapp")
                if shortcut is None:
                    raise ActionError("WhatsApp isn't installed on this PC") from None
                os.startfile(str(shortcut))  # noqa: S606
        time.sleep(0.5)
    raise ActionError("WhatsApp didn't open")


def _click_center(element) -> None:  # noqa: ANN001
    """A real left click in the middle of an element (works where the mouse is; no fail-safe)."""
    rect = element.CurrentBoundingRectangle
    x, y = (rect.left + rect.right) // 2, (rect.top + rect.bottom) // 2
    user32 = ctypes.windll.user32
    user32.SetCursorPos(x, y)
    time.sleep(0.05)
    user32.mouse_event(0x0002, 0, 0, 0, 0)  # left down
    user32.mouse_event(0x0004, 0, 0, 0, 0)  # left up


def _composer(ui: _Uia):  # noqa: ANN202
    """WhatsApp's message box and the chat it's for ("Type a message to <name>"), or None. Opening a chat
    replaces the box, so one found a moment before can be gone when it's read: then look again."""
    for _ in range(3):
        box = ui.first(COMPOSER_PREFIX.strip(), "Edit", contains=True)
        if box is None:
            return None
        try:
            return box, (box.CurrentName or "")[len(COMPOSER_PREFIX):].strip()
        except Exception as exc:
            if not _gone(exc):
                raise
        time.sleep(0.2)
    return None


def _composer_chat(ui: _Uia) -> str | None:
    found = _composer(ui)
    return found[1] if found else None


def _listed(ui: _Uia) -> list[tuple[object, str]] | None:
    """The rows WhatsApp lists now, as (row, chat name); None while it's still redrawing them."""
    try:
        return [(row, title) for row in ui.rows() if (title := _row_title(row.CurrentName or ""))]
    except Exception as exc:
        if _gone(exc):
            return None
        raise


def _search_rows(ui: _Uia, search, query: str, contact: str) -> list[tuple[object, str]]:  # noqa: ANN001
    """Type `query` into WhatsApp's search and read the rows it lists, as (row, chat name): once they stop
    changing, or one is `contact` exactly, for up to 3 seconds."""
    ui.set_value(search, query)
    rows: list[tuple[object, str]] = []
    deadline = time.monotonic() + SEARCH_SECONDS
    while time.monotonic() < deadline:
        time.sleep(0.4)
        now = _listed(ui)
        if now is None:  # a row went while it was being read: the list is still changing
            continue
        if any(_same_person(contact, title) for _row, title in now) or (now and [t for _r, t in now] == [t for _r, t in rows]):
            return now
        rows = now
    return rows


def _open_row(ui: _Uia, row, title: str) -> str | None:  # noqa: ANN001
    """Click a chat in the list; its name as the message box then reports it."""
    _click_center(row)  # what a person does; focusing a row doesn't move WhatsApp's own selection
    chat = None
    for _ in range(15):
        time.sleep(0.2)
        chat = _composer_chat(ui)
        if chat and _same_person(title, chat):
            return chat
    return chat


def _open_chat(ui: _Uia, contact: str, nothing: str = "nothing was sent") -> str | None:
    """Search for `contact` and open its chat: the one named exactly that, else the one whose name clearly fits
    best (name_score) - "Mronal" opens Mrunal, "Abhi Choudhury" opens Abhi Chaudhari. Returns the chat's real
    name as WhatsApp's message box reports it ("Type a message to <name>"), or None. When no chat clearly fits,
    raises naming the closest, so the user can say which. WhatsApp redraws its list as search results come in,
    so a row read a moment before can be gone when it's clicked: then it starts again (nothing is sent or
    called before the chat is open and checked)."""
    for attempt in range(3):
        try:
            return _search_and_open(ui, contact, nothing)
        except Exception as exc:
            if not _gone(exc) or attempt == 2:
                raise
            time.sleep(0.5)
    return None


def _search_and_open(ui: _Uia, contact: str, nothing: str) -> str | None:
    if (_composer_chat(ui) or "") and _same_person(contact, _composer_chat(ui) or ""):
        return _composer_chat(ui)  # already open
    # While it holds a search, the box loses its name; the message box always has one.
    search = ui.first("Search or start a new chat", "Edit") or ui.first("", "Edit")
    if not search:  # WhatsApp is showing Status, Calls or Settings: close any popup, go back to Chats
        _keystroke(VK_ESCAPE)
        chats = ui.first("Chats", "Button")
        if chats is not None:
            ui.press(chats)
        time.sleep(0.8)
        search = ui.first("Search or start a new chat", "Edit") or ui.first("", "Edit")
    if not search:
        raise ActionError("couldn't find WhatsApp's search box (is WhatsApp signed in?)")
    words = [w for w in re.findall(r"\w+", contact) if w.lower() not in _NAME_FILLER]
    # WhatsApp's search wants a name's exact letters, so next the first name alone ("Abhi" lists Abhi Chaudhari
    # for "Abhi Choudhury"), then its first two letters ("Mr" lists Mrunal for "Mronal").
    queries = [contact] + ([words[0]] if len(words) > 1 else []) + ([words[0][:2]] if words and len(words[0]) >= 4 else [])
    close: dict[str, float] = {}
    for query in dict.fromkeys(queries):
        rows = _search_rows(ui, search, query, contact)
        exact = next(((row, title) for row, title in rows if _same_person(contact, title)), None)
        if exact:
            return _open_row(ui, *exact)
        best: dict[str, tuple[float, object]] = {}  # chat name -> (fit, its row); a name can be listed twice
        for row, title in rows:
            score = name_score(contact, title)
            if score > best.get(title, (-1.0, None))[0]:
                best[title] = (score, row)
        ranked = sorted(best.items(), key=lambda item: -item[1][0])
        if ranked and ranked[0][1][0] >= NAME_MATCH and (len(ranked) == 1 or ranked[1][1][0] < ranked[0][1][0] - 0.1):
            return _open_row(ui, ranked[0][1][1], ranked[0][0])
        for title, (score, _row) in ranked:
            close[title] = max(close.get(title, 0.0), score)
    nearest = [title for title, score in sorted(close.items(), key=lambda item: -item[1]) if score >= NAME_CLOSE][:3]
    if nearest:
        raise ActionError(f"no WhatsApp chat called {contact!r} was found; the closest are "
                          + ", ".join(repr(n) for n in nearest) + f" - ask which one they mean; {nothing}")
    return None


def _chat_for(contact: str, nothing: str) -> tuple[int, _Uia, str]:
    """WhatsApp with the chat named `contact` open and checked: (its window, its UI, the chat's real name)."""
    hwnd = _whatsapp_window()
    refocus(hwnd)
    ui = _Uia(hwnd)
    chat = _open_chat(ui, contact, nothing)
    if chat is None:
        raise ActionError(f"no WhatsApp chat called {contact!r} was found; {nothing}")
    if not (_same_person(contact, chat) or name_score(contact, chat) >= NAME_MATCH):
        raise ActionError(f"the chat that opened is {chat!r}, not {contact!r}; {nothing}")
    return hwnd, ui, chat


def _still_open(ui: _Uia, chat: str, doing: str, nothing: str):  # noqa: ANN202
    """The message box, when `chat` is still the one open (after the dialog, the user may have clicked another)."""
    found = _composer(ui)
    if found is None or not _same_person(chat, found[1]):
        raise ActionError(f"the chat changed before {doing}; {nothing}")
    return found[0]


def send_message(contact: str, text: str, app: str = "whatsapp", **_: object) -> dict:
    """Open the chat named `contact` in WhatsApp, check it really is that chat, show the user who
    and what, and send only if they allow it."""
    hwnd, ui, chat = _chat_for(contact, "nothing was sent")
    # The dialog names the chat as WhatsApp does ("Mrunal Patil" for a "Mronal" that was said): that's the OK.
    if confirm_hook is None or not confirm_hook("Send a WhatsApp message", f"To: {chat}\n\n{text[:1500]}"):
        raise ActionError("The user declined this on the PC.")  # nobody to ask = never send
    refocus(hwnd)
    composer = _still_open(ui, chat, "sending", "nothing was sent")
    composer.SetFocus()
    time.sleep(0.15)
    _paste_into_focus(text)
    _keystroke(VK_RETURN)
    time.sleep(0.6)
    return {"sent_to": chat, "app": "WhatsApp", "characters": len(text)}


CALL_SECONDS = 6.0  # how long WhatsApp gets to show its call screen
_HANG_UP = ("End call", "Hang up", "Leave call")


def _call_control(ui: _Uia, video: bool):  # noqa: ANN202
    """The open chat's button for a voice (or video) call. Where WhatsApp has one "Call" button with a menu of
    both, the menu's item: the menu is opened with Expand, which never starts a call by itself."""
    names = ("Video call",) if video else ("Voice call", "Audio call")

    def named():  # noqa: ANN202 - the exact name first: a "Missed voice call" note in the chat isn't the button
        return next((found for contains in (False, True) for name in names for control in ("Button", "MenuItem")
                     if (found := ui.first(name, control, contains=contains)) is not None), None)

    if (found := named()) is not None:
        return found
    menu = ui.first("Call", "Button")
    if menu is None or not ui.expand(menu):
        return None
    for _ in range(8):
        time.sleep(0.2)
        if (found := named()) is not None:
            return found
    _keystroke(VK_ESCAPE)  # close the menu again
    return None


def _call_started(ui: _Uia, before: set[int]) -> bool:
    """Whether WhatsApp shows a call within CALL_SECONDS: a window of its own that wasn't there, or a hang-up button."""
    deadline = time.monotonic() + CALL_SECONDS
    while time.monotonic() < deadline:
        time.sleep(0.5)
        if any(process.lower().startswith("whatsapp") and handle not in before for _t, process, _pid, handle in _windows()):
            return True
        try:
            if any(ui.first(name, "Button", contains=True) is not None for name in _HANG_UP):
                return True
        except Exception as exc:
            if not _gone(exc):  # else WhatsApp is redrawing for the call: look again
                raise
    return False


def call(contact: str, video: bool = False, app: str = "whatsapp", **_: object) -> dict:
    """Open the chat named `contact` in WhatsApp, check it really is that chat, show the user who, and start a
    voice (or video) call only if they allow it."""
    kind = "video call" if video else "voice call"
    hwnd, ui, chat = _chat_for(contact, "no call was made")
    if confirm_hook is None or not confirm_hook(f"Start a WhatsApp {kind}", f"To: {chat}"):
        raise ActionError("The user declined this on the PC.")  # nobody to ask = never call
    refocus(hwnd)
    _still_open(ui, chat, "calling", "no call was made")
    button = _call_control(ui, video)
    if button is None:
        about = sorted({n for n in ui.names("Button") if "call" in n.lower()})[:5]
        raise ActionError(f"WhatsApp shows no {kind} button in the chat with {chat} "
                          + (f"(its buttons about calls: {', '.join(repr(n) for n in about)})" if about
                             else "(a group, or a business that takes no calls?)") + "; no call was made")
    before = {handle for _t, process, _pid, handle in _windows() if process.lower().startswith("whatsapp")}
    ui.press(button)
    if _call_started(ui, before):
        return {"calling": chat, "app": "WhatsApp", "kind": kind}
    return {"calling": chat, "app": "WhatsApp", "kind": kind, "call_screen_seen": False,
            "note": f"WhatsApp's {kind} button was pressed, but its call screen didn't show within "
                    f"{CALL_SECONDS:.0f} seconds: it may still be connecting, or asking for the microphone"}


ACTIONS = {
    "system_info": system_info, "list_files": list_files, "find_files": find_files, "read_file": read_file,
    "list_windows": list_windows, "get_clipboard": get_clipboard, "screenshot": screenshot, "open": open,
    "media": media, "set_volume": set_volume, "lock_screen": lock_screen, "focus_window": focus_window,
    "close_window": close_window, "type_text": type_text, "press_keys": press_keys, "click": click,
    "write_file": write_file, "move_file": move_file, "delete_file": delete_file, "set_clipboard": set_clipboard,
    "go_to": go_to, "search_site": search_site, "play": play, "now_playing": now_playing,
    "set_brightness": set_brightness, "set_theme": set_theme, "open_settings": open_settings, "timer": timer,
    "list_apps": list_apps, "close_app": close_app, "copy_file": copy_file, "create_folder": create_folder,
    "organize_folder": organize_folder, "write_in_app": write_in_app, "write_document": write_document,
    "read_screen": read_screen,
    "click_element": click_element, "fill_field": fill_field, "send_message": send_message, "call": call,
    "do_not_disturb": do_not_disturb, "alarm": alarm, "say": say, "confirm_plan": confirm_plan,
    "notifications": lambda **a: notices.read(int(a.get("hours") or 24), str(a.get("app") or "")),
    # The vision model's eyes and hands (backend app/vision.py): only after you allowed pc_look or pc_click_on.
    "capture": capture, "click_at_checked": click_at_checked,
    # Custodian's own browser, over the DevTools protocol (browser.py): risky clicks and sends ask first.
    "browser_open": lambda **a: browser.browser_open(**a), "browser_read": lambda **a: browser.browser_read(**a),
    "browser_click": lambda **a: browser.browser_click(**a, confirm_hook=confirm_hook),
    "browser_fill": lambda **a: browser.browser_fill(**a, confirm_hook=confirm_hook),
    "browser_download": lambda **a: browser.browser_download(**a), "browser_tabs": lambda **a: browser.browser_tabs(**a),
    # Undo steps: the server sends these only from a record one of the actions above returned.
    "restore_deleted": restore_deleted, "undo_moves": undo_moves, "remove_empty_folder": remove_empty_folder,
    "recycle_created": recycle_created, "restore_backup": restore_backup, "cancel_timer": cancel_timer,
}
SELF_CONFIRMING = {"send_message", "call", "browser_click", "browser_fill"}  # they ask once they know exactly what they'll do
_NOTHING_DONE = {"send_message": "nothing was sent - ask again", "call": "no call was made - ask again"}
# Of those, the ones the server doesn't mark "confirm": they ask only when what they found is risky
# (a submit or "Place order" button, Enter in a form), so following a link never asks.
ASK_WHEN_RISKY = {"browser_click", "browser_fill"}
NEEDS_PROTECTED = {"write_file", "move_file", "delete_file", "copy_file", "create_folder", "organize_folder", "write_document",
                   "restore_deleted", "undo_moves", "remove_empty_folder", "recycle_created", "restore_backup"}


def run(function: str, args: dict, protected: list[str]) -> object:
    action = ACTIONS.get(function)
    if action is None:
        raise ActionError(f"this desktop app doesn't know how to {function}; update it")
    extra = {"protected": protected} if function in NEEDS_PROTECTED else {}
    try:
        return action(**args, **extra)
    except Exception as exc:
        if type(exc).__name__ == "FailSafeException":  # pyautogui's emergency stop
            raise ActionError("stopped: the mouse is in a corner of the screen (Custodian's emergency stop); "
                              "move it away and ask again") from None
        if type(exc).__name__ == "COMError":  # Windows UI Automation's codes, in words
            log.warning("%s: UI Automation error", function, exc_info=True)
            code = exc.args[0] if exc.args and isinstance(exc.args[0], int) else 0
            what = ("the app redrew its window while Custodian was reading it" if _gone(exc)
                    else f"Windows couldn't read or use the app's window (UI Automation error {code & 0xFFFFFFFF:#010x})")
            raise ActionError(f"{what}, so it stopped; {_NOTHING_DONE.get(function, 'ask again')}") from None
        raise


def _window_title(window: str) -> str:
    if window:
        best = best_window(window)
        return best[0] if best else window
    return foreground()[0] or "the window in front"


def foreground() -> tuple[str, int]:
    """(title, handle) of the window in front."""
    user32 = ctypes.windll.user32
    hwnd = user32.GetForegroundWindow()
    length = user32.GetWindowTextLengthW(hwnd)
    buffer = ctypes.create_unicode_buffer(length + 1)
    user32.GetWindowTextW(hwnd, buffer, length + 1)
    return buffer.value, hwnd


def refocus(hwnd: int) -> None:
    """Give the keyboard back to a window (the confirmation dialog took it)."""
    user32 = ctypes.windll.user32
    if hwnd and user32.IsWindow(hwnd):
        if user32.IsIconic(hwnd):
            user32.ShowWindow(hwnd, 9)  # SW_RESTORE: a minimised window has nothing on screen to use
            time.sleep(0.6)
        user32.keybd_event(0x12, 0, 0, 0)  # an Alt tap lets this process move the foreground
        user32.keybd_event(0x12, 0, 2, 0)
        user32.SetForegroundWindow(hwnd)
        time.sleep(0.4)


def describe(function: str, args: dict) -> str:
    """What the confirmation dialog shows: the action in plain words, with its exact details."""
    shown = {
        "type_text": lambda a: f"Type this into the window in front:\n\n{a['text'][:500]}",
        "press_keys": lambda a: f"Press {a['keys']}",
        "click": lambda a: f"Click ({a.get('button', 'left')}) at x={a['x']}, y={a['y']}",
        "write_file": lambda a: f"{'Overwrite' if a.get('overwrite') else 'Create'} the file\n{a['path']}\n\n{a['content'][:400]}",
        "move_file": lambda a: f"Move\n{a['source']}\nto\n{a['destination']}",
        "delete_file": lambda a: f"Move to the Recycle Bin:\n{a['path']}",
        "close_window": lambda a: f"Close the window \u201c{(best_window(a['title']) or [a['title']])[0]}\u201d (unsaved work may be lost)",
        "read_file": lambda a: f"Read the file\n{a['path']}\nand send its text to Custodian",
        "get_clipboard": lambda a: "Send the text on your clipboard to Custodian",
        "notifications": lambda a: (f"Send your {a.get('app') + ' ' if a.get('app') else ''}notifications from the last "
                                    f"{a.get('hours', 24)} hours (who wrote, and what), and the unread counts, to Custodian"),
        "screenshot": lambda a: "Take a screenshot and send it to Custodian to describe",
        "close_app": lambda a: f"Quit {app_process(a['name']) or a['name']} and close all its windows (it may ask to save)",
        "organize_folder": lambda a: f"Sort the loose files in\n{a['path']}\ninto subfolders by type (Images, Documents, ...)",
        "write_in_app": lambda a: f"Open {a.get('app', 'notepad')} and write:\n\n{a['text'][:500]}",
        "click_element": lambda a: f"Click \u201c{a['name']}\u201d in \u201c{_window_title(a.get('window', ''))}\u201d",
        "call": lambda a: f"Start a WhatsApp {'video' if a.get('video') else 'voice'} call to {a['contact']}",
        "fill_field": lambda a: f"Type into \u201c{a['name']}\u201d in \u201c{_window_title(a.get('window', ''))}\u201d:\n\n{a['text'][:300]}",
        "confirm_plan": lambda a: a.get("details") or "\n".join(f"\u2022 {step}" for step in a["steps"][:12]),
        "browser_download": lambda a: "Download into your Downloads folder:\n"
        + (a.get("url") or f"the link “{a.get('name') or a.get('ref')}” on the page in Custodian's browser"),
        "capture": lambda a: (f"Look at your screen to {a['purpose']}" if a.get("purpose") else "Look at your screen")
        + "\n\nA screenshot goes to Custodian's vision model to be read, and isn't kept.",
    }
    return shown.get(function, lambda a: f"{function}: {a}")(args)
