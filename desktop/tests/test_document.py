"""pc_write_document: a real Word file with headings, bullets and the sources as links, always a new
file, never inside a protected folder."""

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

from custodian_desktop import actions  # noqa: E402

SECTIONS = [{"heading": "Summary", "text": "Get the Vivobook 16 [1].\n\nIt has the longest battery life."},
            {"heading": "Options", "text": "- Vivobook 16: Rs 62,990 [1]\n- IdeaPad Slim 5: Rs 71,500 [2]"}]
SOURCES = [{"title": "Laptop review", "url": "https://reviews.example/vivobook"},
           {"title": "Buying guide", "url": "https://guide.example/laptops"}]


@pytest.fixture
def documents(tmp_path, monkeypatch):
    opened = []
    monkeypatch.setattr(actions, "known_folder", lambda name: tmp_path if name == "documents" else None)
    monkeypatch.setattr(actions.os, "startfile", opened.append, raising=False)
    return tmp_path, opened


def test_a_report_becomes_a_word_document_with_links_and_opens(documents):
    folder, opened = documents
    result = actions.run("write_document", {"title": "Laptops under 80000", "sections": SECTIONS, "sources": SOURCES}, [])
    path = folder / "Laptops under 80000.docx"
    assert result["saved"] == str(path) and opened == [str(path)]
    assert result["undo"] == {"function": "recycle_created", "args": {"path": str(path)}, "label": "delete Laptops under 80000.docx"}
    doc = Document(str(path))
    lines = [(p.style.name, p.text) for p in doc.paragraphs]
    assert lines[0] == ("Title", "Laptops under 80000")
    assert ("Heading 1", "Options") in lines and ("List Bullet", "IdeaPad Slim 5: Rs 71,500 [2]") in lines
    assert ("Normal", "It has the longest battery life.") in lines
    assert lines[-1][1].startswith("[2] Buying guide - https://guide.example/laptops")
    links = {r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")}
    assert links == {"https://reviews.example/vivobook", "https://guide.example/laptops"}


def test_never_over_another_file_and_never_in_a_protected_folder(documents):
    folder, _ = documents
    args = {"title": "Plan: Q3/Q4?", "sections": SECTIONS}
    first = actions.run("write_document", args, [])
    second = actions.run("write_document", args, [])
    assert first["saved"].endswith("Plan Q3Q4.docx") and second["saved"].endswith("Plan Q3Q4 (2).docx")
    with pytest.raises(actions.ActionError, match="never changes"):
        actions.run("write_document", args, [str(folder)])


def test_a_report_given_as_one_text_becomes_headed_sections(documents):
    folder, _ = documents
    body = "For most people, the Vivobook [1].\n\n## Top picks\n- **Vivobook 16**: Rs 62,990 [1]\n- IdeaPad: Rs 71,500\n\n## What to choose\nThe Vivobook."
    actions.run("write_document", {"title": "Laptops", "body": body, "sources": SOURCES[:1]}, [])
    lines = [(p.style.name, p.text) for p in Document(str(folder / "Laptops.docx")).paragraphs]
    assert ("Normal", "For most people, the Vivobook [1].") in lines  # before the first heading: no heading
    assert ("Heading 1", "Top picks") in lines and ("List Bullet", "Vivobook 16: Rs 62,990 [1]") in lines  # no ** left
    assert ("Heading 1", "What to choose") in lines


def test_the_title_and_a_written_sources_list_arent_repeated(documents):
    folder, _ = documents
    body = "# Laptops under 80000\nGet the Vivobook [1].\n\n## Sources\n1. Laptop review"
    actions.run("write_document", {"title": "Laptops under 80,000", "body": body, "sources": SOURCES[:1]}, [])
    lines = [(p.style.name, p.text) for p in Document(str(folder / "Laptops under 80,000.docx")).paragraphs]
    assert [t for s, t in lines if s.startswith("Heading")] == ["Sources"]  # only the linked list, and no title twice
    assert ("Normal", "Get the Vivobook [1].") in lines
